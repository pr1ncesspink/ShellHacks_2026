"""Shared project normalization and workbook identity rules; no document extraction."""

from __future__ import annotations

import re
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook

DESC = "Dominion Energy South Carolina"
GPC = "Georgia Power"
HEADERS = ["overlap_id", "distance_mi", "time_gap (day)", "utility_a",
           "project_id_a", "project_name_a", "utility_b", "project_id_b", "project_name_b"]
CEII = re.compile(r"critical\s+energy\s+infrastructure\s+information|confidential\s+CEII|"
                  r"CEII\s*[-–:]?\s*confidential|(?:CFR|C\.?F\.?R\.?)\s*(?:sec\.?|section|§)?\s*388\.113", re.I)


def normalize(text):
    text = str(text or "").translate(str.maketrans({"–": "-", "—": "-", "−": "-", "’": "'"}))
    text = re.sub(r"(?<=\d)\s*kv\b", " kV", text, flags=re.I)
    return re.sub(r"\s+", " ", re.sub(r"\bkv\b", "kV", text, flags=re.I)).strip()


def name_key(text):
    return re.sub(r"[^a-z0-9]", "", normalize(text).lower())


def endpoint_key(text):
    text = normalize(text).lower()
    text = re.sub(r"\bsub(?:station)?\.?\b", "", text)
    text = re.sub(r"\bjct\b", "junction", text)
    text = re.sub(r"\bft\b", "fort", text)
    return name_key(text)


def iso_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def read_seed(path):
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if list(next(wb["overlaps"].iter_rows(values_only=True))) != HEADERS:
            raise ValueError("The workbook overlaps headers do not match the required nine-column schema")
        rows = wb["projects"].iter_rows(values_only=True)
        columns = list(next(rows))
        required = {"project_id", "utility", "state", "project_name", "name_a", "name_b",
                    "lat_a", "lon_a", "lat_b", "lon_b", "in_service_date"}
        if not required.issubset(columns):
            raise ValueError(f"Missing workbook project fields: {sorted(required - set(columns))}")
        result = []
        for row in rows:
            if not any(v is not None for v in row):
                continue
            p = dict(zip(columns, row))
            p = {k: p.get(k) for k in required}
            p.update(source_document=Path(path).name, source_page="", raw_text="", warnings=[],
                     published_project_id="", description="", need="", status="",
                     dates=[], annual_costs={}, voltages_kv=[], line_length_mi=None,
                     asset_type="", total_cost=None, previous_cost=None,
                     location_evidence={}, confidence_a=0., confidence_b=0.)
            p["in_service_date"] = iso_date(p["in_service_date"])
            p["dates"] = [p["in_service_date"]] if p["in_service_date"] else []
            for side in ("a", "b"):
                lat, lon = p[f"lat_{side}"], p[f"lon_{side}"]
                if isinstance(lat, (int, float)) and isinstance(lon, (int, float)) and -90 <= lat <= 90 and -180 <= lon <= 180:
                    p[f"confidence_{side}"] = 1.
                    p["location_evidence"][side] = {"method": "sponsor_workbook", "confidence": 1.}
                else:
                    p[f"lat_{side}"] = p[f"lon_{side}"] = None
            result.append(p)
        if len({p["project_id"] for p in result}) != len(result):
            raise ValueError("Duplicate project IDs in starter workbook")
        return result
    finally:
        wb.close()


def endpoint_names(title):
    """Conservative endpoint extraction: complex scopes require review, not guessed centroids."""
    title = normalize(title)
    title = re.sub(r"^(?:SAV|AUG|ATL):\s*", "", title, flags=re.I)
    prefix = re.split(r"\b\d+(?:\.\d+)?\s*(?:-\s*\d+(?:\.\d+)?\s*)?kV\b|:", title, maxsplit=1, flags=re.I)[0]
    # A tap is an asset label, not an arbitrary split within a proper name.
    prefix = re.sub(r"\s+(?:transmission\s+)?tap\b.*$", "", prefix, flags=re.I)
    prefix = re.sub(r"\([^)]*\)", "", prefix).strip()
    parts = [x.strip() for x in re.split(r"\s*-\s*|\s+to\s+", prefix, flags=re.I) if x.strip()]
    if len(parts) > 2 or "&" in prefix or re.search(r"\band\b", prefix, re.I):
        return "", "", ["complex_endpoint_scope: manual endpoint review required"]
    if not parts:
        return "", "", ["no_named_endpoints"]
    return parts[0], parts[1] if len(parts) == 2 else "", []


def scope_endpoints(title, description):
    """Prefer an explicitly described sub-section over endpoints of its parent line."""
    title, description = normalize(title), normalize(description)
    # Several separate substations or differently named parallel lines need review.
    if re.search(r"\bsub(?:station)?\b.*\bsub(?:station)?\b.*\bsub(?:station)?\b", title, re.I):
        return "", "", ["complex_endpoint_scope: multiple substations"]
    if re.search(r"kV\s*&\s*[A-Za-z]", title, re.I):
        return "", "", ["complex_endpoint_scope: multiple lines"]
    section = re.search(r":\s*(?:Rebuild\s+)?(.+?)\s+(?:Transmission\s+)?Section\b", title, re.I)
    if section and re.search(r"-|\bto\b", section.group(1), re.I):
        return endpoint_names(section.group(1))
    # Titles sometimes name the parent corridor then a specific rebuild segment.
    section = re.search(r":\s*Rebuild Line from\s+(.+)$", title, re.I)
    if section:
        return endpoint_names(re.sub(r"\s+Tie$", "", section.group(1), flags=re.I))
    section = re.search(r"kV:\s*([^:]+?\s+-\s+[^:]+)$", title, re.I)
    if section:
        return endpoint_names(section.group(1))
    route = re.search(r"\bfrom\s+(.+?)\s+to\s+(.+?)(?:\s+to\s+(?:provide|feed)\b|[.;]|$)", description, re.I)
    if route:
        a, b = route.groups()
        a = re.sub(r"^the\s+", "", a, flags=re.I)
        b = re.sub(r"\s+\d+(?:/\d+)?\s*kV.*$", "", b, flags=re.I)
        b = re.sub(r"\s+Substation$", "", b, flags=re.I)
        # Exclude conductor-upgrade grammar (from 397.5 ACSR ... to 1272 ACSR).
        if not re.match(r"\d", a) and not re.match(r"\d", b) and len(a.split()) <= 6 and len(b.split()) <= 6:
            return a, b, []
    route = re.search(r"^Construct\s+(.+?)\s+\d+(?:/\d+)?\s*kV", description, re.I)
    if route and "-" in route.group(1):
        return endpoint_names(route.group(1))
    return endpoint_names(title)


def merge_seed(parsed, seed):
    remaining = {p["project_id"]: p for p in seed}
    result = []
    for p in parsed:
        candidates = sorted(((SequenceMatcher(None, name_key(p["project_name"]), name_key(s["project_name"])).ratio(), s)
                             for s in remaining.values() if s["utility"] == p["utility"]),
                            key=lambda x: (-x[0], x[1]["project_id"]))
        if candidates and candidates[0][0] >= .94 and (len(candidates) == 1 or candidates[0][0] - candidates[1][0] >= .05):
            s = candidates[0][1]
            for key in ("project_id", "name_a", "name_b", "lat_a", "lon_a", "lat_b", "lon_b",
                        "confidence_a", "confidence_b", "location_evidence"):
                p[key] = s[key]
            p["seed_project_name"] = s["project_name"]
            p["warnings"] = [w for w in p["warnings"] if not w.startswith("complex_endpoint_scope")]
            del remaining[s["project_id"]]
        result.append(p)
    return result + list(remaining.values())
