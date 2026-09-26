"""Compose document extraction with the existing identity, geography and overlap rules."""

from __future__ import annotations

from contextlib import ExitStack
import csv
import hashlib
import json
import os
from pathlib import Path
import tempfile

from openpyxl import load_workbook

from .collisions import collisions, write_csv
from .locations import Resolver, center
from .records import HEADERS, iso_date, merge_seed, read_seed
from .config import SnowflakeSettings
from .extraction import Project, SCHEMA_VERSION, deduplicate, extract_document, validate_document
from .snowflake import SnowflakeClient


OUTPUTS = ("projects.json", "overlaps.json", "workbook.json", "extraction_audit.json",
           "pipeline_audit.json", "location_review.json", "collisions.csv")
WORKBOOK_FIELDS = ("project_id", "utility", "state", "project_name", "name_a", "lat_a", "lon_a",
                   "name_b", "lat_b", "lon_b", "lat_center", "lon_center", "in_service_date")


def structured_projects(path: Path):
    """Read the sponsor workbook or an equivalent projects CSV without AI reinterpretation."""
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if path.suffix.lower() == ".xlsx":
        rows = read_seed(path)
    else:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            required = {"project_id", "project_name", "utility"}
            if not required.issubset(reader.fieldnames or []):
                raise ValueError("Projects CSV requires project_id, project_name, and utility columns")
            rows = list(reader)
    projects = []
    for index, row in enumerate(rows, 2):
        values = {key: (None if value == "" else value) for key, value in row.items()
                  if key in Project.model_fields}
        # Spreadsheet formula centers and overlap columns are intentionally not imported.
        values.pop("lat_center", None)
        values.pop("lon_center", None)
        original_date = row.get("in_service_date")
        parsed = iso_date(original_date) or None
        warnings = list(values.get("warnings") or [])
        if original_date and not parsed:
            warnings.append("missing_or_invalid_date")
        values.update(in_service_date=parsed, dates=[parsed] if parsed else [],
                      source_document=path.name, source_page=None, raw_text=values.get("raw_text") or "",
                      source_references=[{"document": path.name, "sha256": digest,
                                          "sheet": "projects" if path.suffix.lower() == ".xlsx" else None,
                                          "row": index}], warnings=warnings)
        project = Project.model_validate(values).model_dump(mode="json")
        for side in ("a", "b"):
            if project[f"lat_{side}"] is not None and project[f"lon_{side}"] is not None:
                project[f"confidence_{side}"] = 1.0
                project["location_evidence"][side] = {"method": "structured_input", "confidence": 1.0}
            else:
                project[f"lat_{side}"] = project[f"lon_{side}"] = None
                project[f"confidence_{side}"] = 0.0
        projects.append(project)
    if len({p["project_id"] for p in projects}) != len(projects):
        raise ValueError(f"Duplicate project IDs in {path.name}")
    return projects, {"file": path.name, "sha256": digest, "parser": "structured_import",
                      "schema_version": SCHEMA_VERSION, "project_count": len(projects)}


def workbook_overlap_ids(paths: list[Path]) -> dict:
    known, ids = {}, {}
    for path in paths:
        if path.suffix.lower() != ".xlsx":
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook["overlaps"].iter_rows(values_only=True)
            headers = next(rows)
            for row in rows:
                if not any(value is not None for value in row):
                    continue
                values = dict(zip(headers, row))
                pair = tuple(sorted((values["project_id_a"], values["project_id_b"])))
                identifier = str(values["overlap_id"])
                if (pair in known and known[pair] != identifier) or (identifier in ids and ids[identifier] != pair):
                    raise ValueError("Conflicting overlap IDs in input workbooks")
                known[pair], ids[identifier] = identifier, pair
        finally:
            workbook.close()
    return known


def build_workbook(projects: list[dict], overlaps: list[dict]) -> dict:
    links = {project["project_id"]: [] for project in projects}
    for row in overlaps:
        a, b = row["project_id_a"], row["project_id_b"]
        links[a].append(b)
        links[b].append(a)
    width = max(3, max((len(values) for values in links.values()), default=0))
    rows = []
    for project in projects:
        linked = sorted(links[project["project_id"]])
        row = {key: project.get(key) for key in WORKBOOK_FIELDS}
        row["overlap_count"] = len(linked)
        row.update({f"overlap_{i + 1}": linked[i] if i < len(linked) else None for i in range(width)})
        rows.append(row)
    return {"projects": rows, "overlaps": overlaps}


def run_pipeline(inputs, output_dir, *, starter_workbook=None, utility="", state="", client=None,
                 cache_dir=".cache/geocoding", osm_snapshot=None, refresh_locations=False,
                 user_agent="", audit_inputs=()):
    paths = list(dict.fromkeys(Path(path).resolve() for path in inputs))
    audits = [Path(path).resolve() for path in audit_inputs]
    seed_path = Path(starter_workbook).resolve() if starter_workbook else None
    if not paths:
        raise ValueError("Provide at least one --input file")
    sources = paths + audits + ([seed_path] if seed_path else [])
    if osm_snapshot:
        sources.append(Path(osm_snapshot).resolve())
    out = Path(output_dir).resolve()
    if any((out / filename).resolve() in sources for filename in OUTPUTS):
        raise ValueError("An output path aliases a source file")
    if osm_snapshot and not Path(osm_snapshot).is_file():
        raise ValueError("OSM snapshot must exist as a file")
    if refresh_locations and not user_agent:
        raise ValueError("Network geocoding requires --user-agent")
    # Reject bad inputs before staging anything or spending Cortex credits.
    for path in paths + audits:
        if path.suffix.lower() in {".xlsx", ".csv"} and path not in audits:
            if not path.is_file():
                raise ValueError(f"Input file does not exist: {path.name}")
        else:
            validate_document(path)
    seed, seed_report = structured_projects(seed_path) if seed_path else ([], None)
    known_ids = workbook_overlap_ids(paths + ([seed_path] if seed_path else []))
    projects, source_reports = [], []
    with ExitStack() as stack:
        for path, audit_only in [(path, False) for path in paths] + [(path, True) for path in audits]:
            if path.suffix.lower() in {".xlsx", ".csv"}:
                parsed, report = structured_projects(path)
            else:
                if client is None:
                    client = stack.enter_context(SnowflakeClient(SnowflakeSettings.from_env()))
                parsed, report = extract_document(path, client, utility=utility, state=state, audit_only=audit_only)
            projects.extend(parsed)
            source_reports.append(report)
    if not projects:
        raise ValueError("No construction projects extracted; inspect the document content and extraction schema")
    projects = deduplicate(projects)
    # Existing sponsor identity matching is reused; unmatched seed records are retained.
    projects = merge_seed(projects, seed)
    if len({p["project_id"] for p in projects}) != len(projects):
        raise ValueError("Duplicate project identities after seed matching; review inputs")
    if seed_report:
        source_reports.append(seed_report)
    resolver = Resolver(cache_dir, seed, network=refresh_locations, user_agent=user_agent, snapshot=osm_snapshot)
    resolvable = [p for p in projects if p.get("utility") and p.get("state")]
    review = resolver.resolve(resolvable)
    for project in projects:
        project["lat_center"], project["lon_center"] = center(project)
    comparable = [p for p in projects if p.get("utility")]
    rows, excluded = collisions(comparable)
    excluded.extend({"project_id": p["project_id"], "reason": "missing_utility"}
                    for p in projects if not p.get("utility"))
    for row in rows:
        pair = tuple(sorted((row["project_id_a"], row["project_id_b"])))
        row["overlap_id"] = known_ids.get(pair) or "OVL_" + hashlib.sha256("|".join(pair).encode()).hexdigest()[:20]
        row["distance_mi"] = float(row["distance_mi"])
    projects = [Project.model_validate(p).model_dump(mode="json") for p in sorted(projects, key=lambda p: p["project_id"])]
    workbook = build_workbook(projects, rows)
    audit = {"schema_version": SCHEMA_VERSION, "project_count": len(projects), "overlap_count": len(rows),
             "eligible_overlap_ids": [row["overlap_id"] for row in rows if row["time_gap (day)"] <= 365],
             "collision_exclusions": excluded, "geocoding": resolver.stats,
             "rules": {"distance_miles_exclusive": 25, "max_in_service_gap_days_inclusive": 365,
                       "timing_basis": "in_service_date_proxy", "overlaps_export": "all_geographic_candidates"},
             "warnings": [{"project_id": p["project_id"], "warnings": p["warnings"]} for p in projects if p["warnings"]],
             "attribution": "OSM-derived matches: © OpenStreetMap contributors, https://www.openstreetmap.org/copyright"}
    artifacts = {"projects.json": projects, "overlaps.json": rows, "workbook.json": workbook,
                 "extraction_audit.json": {"sources": source_reports}, "pipeline_audit.json": audit,
                 "location_review.json": review}
    # Finish extraction/validation and serialize every artifact before replacing prior outputs.
    serialized = {name: json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n"
                  for name, value in artifacts.items()}
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".document-export-", dir=out) as directory:
        pending = Path(directory)
        for name, content in serialized.items():
            (pending / name).write_text(content, encoding="utf-8")
        csv_rows = [{**row, "distance_mi": f"{row['distance_mi']:.2f}"} for row in rows]
        write_csv(pending / "collisions.csv", HEADERS, csv_rows)
        for filename in OUTPUTS:
            os.replace(pending / filename, out / filename)
    return audit
